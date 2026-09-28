"""Import new clients and update nonblank fields matched by registration number."""

import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models.functions import Lower, Trim
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .models import Client
from .profile_fields import PROFILE_FIELDS


def header_key(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


ALIASES = {
    "birthdate": "date_of_birth",
    "nextofkin": "next_of_kin_name",
    "regdate": "registration_date",
    "recommendedby": "referred_by",
    "savegoal": "savings_goal",
    "highestlevel": "highest_education_level",
    "nextofkinrelationship": "next_of_kin_relationship",
    "savingsfreq": "savings_frequency",
    "minsavings": "minimum_savings_amount",
    "grpmembers": "group_member_count",
    "grpobjective": "group_objective",
    "grpsavingcyclelength": "group_savings_cycle_months",
    "grpchair": "group_chairperson",
    "grpsecreatry": "group_secretary",
    "grpsecretary": "group_secretary",
    "grptreasurer": "group_treasurer",
    "profession": "occupation",
    "typeofwork": "occupation",
}
DATES = {"date_of_birth", "registration_date"}
NUMBERS = {"group_member_count", "group_savings_cycle_months"}


def clean_value(field, value, epoch, legacy_cycle=False):
    if value is None or str(value).strip().lower() in {"", "n/a", "na", "none", "null"}:
        return None if field in DATES | NUMBERS | {"minimum_savings_amount"} else ""
    if field in DATES:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, (int, float)):
            converted = from_excel(value, epoch)
            if not isinstance(converted, datetime):
                raise ValueError(f"{field}: invalid Excel date")
            return converted.date()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(str(value).strip(), fmt).date()
            except ValueError:
                pass
        raise ValueError(f"{field}: use YYYY-MM-DD or DD/MM/YYYY")
    value = str(value).strip()
    if field == "gender":
        return {"m": "Male", "male": "Male", "f": "Female", "female": "Female"}.get(value.lower(), value)
    if field in {"client_type", "savings_frequency"}:
        return value.lower()
    if field in NUMBERS:
        if field == "group_savings_cycle_months" and legacy_cycle:
            match = re.fullmatch(r"(\d+)\s*(years?|months?)", value.lower())
            if not match:
                raise ValueError("Group savings cycle needs an explicit unit, e.g. 1 year or 12 months")
            return int(match[1]) * (12 if match[2].startswith("year") else 1)
        number = Decimal(value)
        if number != number.to_integral_value():
            raise ValueError(f"{field}: use a whole number")
        return int(number)
    if field == "minimum_savings_amount":
        return Decimal(value.replace(",", ""))
    return value


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    errors: list = field(default_factory=list)


def import_clients(excel_file, *, report=False):
    result = ImportResult()
    errors = result.errors
    workbook = load_workbook(excel_file, read_only=True, data_only=False)
    try:
        sheet = workbook["Updates"] if "Updates" in workbook.sheetnames else workbook.active
        source = sheet.iter_rows(values_only=True)
        headers = [header_key(value) for value in next(source, ())]
        canonical = {header_key(field): field for field in (*PROFILE_FIELDS, "picture")}
        mapping = [canonical.get(key) or ALIASES.get(key) for key in headers]
        if "full_name" not in mapping or "reg_number" not in mapping:
            errors.append("The workbook must contain full_name and reg_number headers.")
            return result if report else errors
        if len([h for h in headers if h]) != len(set(h for h in headers if h)):
            errors.append("Duplicate column headers. Keep one column per field.")
            return result if report else errors
        unknown = [h for h, field in zip(headers, mapping) if h and not field]
        if unknown:
            errors.append("Unknown columns: " + ", ".join(unknown))
            return result if report else errors
        rows = list(source)
        reg_index = mapping.index("reg_number")
        counts = Counter(str(row[reg_index] or "").strip().lower() for row in rows)
        for row_number, row in enumerate(rows, 2):
            if not any(value not in (None, "") for value in row):
                continue
            try:
                reg_key = str(row[reg_index] or "").strip().lower()
                if not reg_key:
                    raise ValueError("reg_number is required")
                if counts[reg_key] > 1:
                    raise ValueError("Duplicate Reg. No in workbook; keep one row per client")
                if any(isinstance(value, str) and value.startswith("=") for value in row):
                    raise ValueError("Replace formulas with their values before importing")
                data = {}
                for header, field, raw in zip(headers, mapping, row):
                    if not field:
                        continue
                    value = clean_value(field, raw, workbook.epoch, header == "grpsavingcyclelength")
                    # Old work columns mix occupations, work sectors, and employment statuses.
                    if field == "occupation" and value:
                        key = value.lower()
                        if key in {"formal", "informal", "formal sector", "informal sector"}:
                            field, value = "employment_sector", key.replace(" sector", "").capitalize()
                        elif key in {"student", "child", "unemployed", "retired"}:
                            field, value = "employment_status", key.capitalize()
                    if value in (None, ""):
                        continue
                    if field in data and str(data[field]).casefold() != str(value).casefold():
                        raise ValueError(f"Conflicting {field} values; use the canonical column names to clarify")
                    data[field] = value
                gender_code = str(data.get("gender", "")).upper()
                account_type = {"G": "group", "C": "group", "J": "joint", "JT": "joint", "M/F": "joint"}.get(
                    gender_code
                )
                if account_type:
                    if data.get("client_type", account_type) != account_type:
                        raise ValueError("Gender account code conflicts with client_type")
                    data["client_type"] = account_type
                    data.pop("gender")
                if data.get("client_type") in {"group", "joint", "organization"}:
                    if data.get("gender"):
                        raise ValueError("Use individual genders only for individual clients")
                    data["gender"] = ""
                kin = data.get("next_of_kin_name", "")
                match = re.fullmatch(r"(.+?)\s*[-/]\s*(\+?\d[\d ]{8,})", kin)
                if match and not data.get("next_of_kin_phone"):
                    data["next_of_kin_name"], data["next_of_kin_phone"] = match[1].strip(), match[2].replace(" ", "")
                with transaction.atomic():
                    matches = list(
                        Client.objects.select_for_update()
                        .annotate(
                            import_reg=Lower(Trim("reg_number")),
                        )
                        .filter(import_reg=reg_key)[:2]
                    )
                    if len(matches) > 1:
                        raise ValueError("Multiple existing clients have this Reg. No; resolve duplicates first")
                    client = matches[0] if matches else Client()
                    if not matches and not data.get("full_name"):
                        raise ValueError("full_name is required for new clients")
                    changed = False
                    for name, value in data.items():
                        if matches and name == "reg_number":
                            continue
                        if str(getattr(client, name)) != str(value):
                            setattr(client, name, value)
                            changed = True
                    client.full_clean(exclude=["picture"])
                    if not matches:
                        client.save()
                        result.created += 1
                    elif changed:
                        client.save()
                        result.updated += 1
                    else:
                        result.unchanged += 1
            except (ValidationError, ValueError, InvalidOperation) as exc:
                errors.append(f"Row {row_number}: {exc}")
    finally:
        workbook.close()
    return result if report else errors
